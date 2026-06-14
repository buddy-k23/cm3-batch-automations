"""Unit tests for the EC-S3 source YAML emitter.

Covers (10 cases):

  * test_shaw_workbook_emits_equivalent_source_yaml — round-trip vs
    committed config/e2e/sources/SHAW.yml (semantic equivalence)
  * test_emitted_yaml_omits_default_boilerplate — EA-S1 / EA-S3 contract
  * test_emitted_yaml_preserves_non_default_tolerance_ignore_fields —
    surgical strip preserved
  * test_emitted_yaml_never_writes_multi_record_keys — EB-S1 / ADR 0005
  * test_emitted_yaml_mapping_path_extension_drives_multi_record —
    EB-S1 dispatch inference
  * test_emitted_yaml_rules_path_empty_for_umbrella — SHAW convention
  * test_gates_block_emitted_from_source_sheet — gates ordering + values
  * test_staging_tables_derived_from_input_files — dedup + order
  * test_emitter_validates_through_SourceConfig — EmitterError on
    invalid output
  * test_emitted_yaml_passes_eb_s2_guardrail — emitted YAML clears the
    EB-S2 banned-keys scanner
"""

from __future__ import annotations

from pathlib import Path

import pytest
import yaml
from openpyxl import Workbook

from src.onboarding.emitters import EmitterError
from src.onboarding.emitters.source_yaml_emitter import (
    SourceYamlEmitter,
    emit_source_yaml,
)
from src.onboarding.workbook_reader import read_workbook
from src.onboarding.workbook_schema import (
    INPUT_FILES_REQUIRED_COLUMNS,
    OUTPUT_FILES_REQUIRED_COLUMNS,
    SOURCE_SHEET_REQUIRED_COLUMNS,
)
from src.pipeline.etl_config import SourceConfig

REPO_ROOT = Path(__file__).resolve().parent.parent.parent
TEMPLATES_DIR = REPO_ROOT / "templates"
SHAW_WORKBOOK = TEMPLATES_DIR / "SHAW_onboarding.xlsx"
COMMITTED_SHAW_YAML = REPO_ROOT / "config" / "e2e" / "sources" / "SHAW.yml"


# ---------------------------------------------------------------------------
# Synthetic-workbook helper (mirrors the EC-S2 test helper, kept local so
# the EC-S3 tests do not couple to the EC-S2 test module).
# ---------------------------------------------------------------------------


def _write_row(ws, row_number: int, values: list[object]) -> None:
    """Write ``values`` into row ``row_number`` of ``ws`` (1-indexed columns)."""
    for col_idx, value in enumerate(values, start=1):
        ws.cell(row=row_number, column=col_idx, value=value)


# Default Source data row used when individual tests don't override cells.
# Mirrors the EC-S2 test helper exactly.
_DEFAULT_SOURCE_DATA: dict[str, object] = {
    "source_code": "TEST",
    "schema_version": 1,
    "release_tag": "2026.M06",
    "description": "Synthetic test workbook.",
    "staging_schema": "APP_INT",
    "output_root": "/tmp/test",
    "java_load_script": "/tmp/load.sh",
    "java_generate_script": "",
    "gate_load_blocking": "true",
    "gate_load_invoke_java": "false",
    "gate_f2s_blocking": "true",
    "gate_generate_blocking": "false",
    "gate_generate_invoke_java": "false",
    "gate_l1_blocking": "true",
    "gate_l2b_blocking": "true",
    "gate_l3_blocking": "false",
    "gate_mr_report_blocking": "false",
}


def _build_workbook(
    tmp_path: Path,
    *,
    source_overrides: dict[str, object] | None = None,
    input_files_rows: list[dict[str, object]] | None = None,
    output_files_rows: list[dict[str, object]] | None = None,
    filename: str = "wb.xlsx",
) -> Path:
    """Synthesise a minimal valid-shape workbook.

    Args:
        tmp_path: pytest tmp_path fixture.
        source_overrides: Per-column overrides for the single Source row.
        input_files_rows: Data rows for InputFiles (header always written).
        output_files_rows: Data rows for OutputFiles.
        filename: Output filename in tmp_path.

    Returns:
        Path to the saved workbook.
    """
    source_overrides = source_overrides or {}
    input_files_rows = input_files_rows or []
    output_files_rows = output_files_rows or []

    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("Source")
    _write_row(ws, 1, list(SOURCE_SHEET_REQUIRED_COLUMNS))
    data_row: list[object] = []
    for col in SOURCE_SHEET_REQUIRED_COLUMNS:
        value = (
            source_overrides[col]
            if col in source_overrides
            else _DEFAULT_SOURCE_DATA[col]
        )
        data_row.append(value)
    _write_row(ws, 2, data_row)

    ws = wb.create_sheet("InputFiles")
    _write_row(ws, 1, list(INPUT_FILES_REQUIRED_COLUMNS))
    for idx, row_dict in enumerate(input_files_rows, start=2):
        row_values = [row_dict.get(col) for col in INPUT_FILES_REQUIRED_COLUMNS]
        _write_row(ws, idx, row_values)

    ws = wb.create_sheet("OutputFiles")
    _write_row(ws, 1, list(OUTPUT_FILES_REQUIRED_COLUMNS))
    for idx, row_dict in enumerate(output_files_rows, start=2):
        row_values = [row_dict.get(col) for col in OUTPUT_FILES_REQUIRED_COLUMNS]
        _write_row(ws, idx, row_values)

    out_path = tmp_path / filename
    wb.save(str(out_path))
    return out_path


# ---------------------------------------------------------------------------
# 1. SHAW round-trip — the central proof-point of the BA workflow.
# ---------------------------------------------------------------------------


def test_shaw_workbook_emits_equivalent_source_yaml():
    """The SHAW workbook round-trips to a YAML semantically equal to the
    committed ``config/e2e/sources/SHAW.yml`` and passes ``SourceConfig``.

    The contract is semantic equivalence via ``yaml.safe_load``, not byte
    equality -- comment blocks and quoting style legitimately diverge.
    """
    workbook = read_workbook(SHAW_WORKBOOK)
    emitted = emit_source_yaml(workbook)

    emitted_data = yaml.safe_load(emitted)
    committed_data = yaml.safe_load(COMMITTED_SHAW_YAML.read_text(encoding="utf-8"))

    assert emitted_data == committed_data, (
        f"Semantic divergence between emitted and committed SHAW.yml.\n"
        f"emitted keys:   {list(emitted_data)}\n"
        f"committed keys: {list(committed_data)}"
    )

    # SourceConfig acceptance -- the emitted YAML is consumable by the
    # downstream Pydantic model, same as the committed file.
    emitted_config = SourceConfig.model_validate(emitted_data)
    committed_config = SourceConfig.model_validate(committed_data)
    assert emitted_config.model_dump() == committed_config.model_dump()


# ---------------------------------------------------------------------------
# 2. Default boilerplate is omitted (EA-S1 / EA-S3 contract).
# ---------------------------------------------------------------------------


def test_emitted_yaml_omits_default_boilerplate(tmp_path):
    """Synthetic workbook with all defaults -- emitted output entries
    contain no ``strict_*`` / ``tolerance`` / ``thresholds`` keys."""
    path = _build_workbook(
        tmp_path,
        input_files_rows=[
            {
                "file_type": "EXAMPLE_IN",
                "glob": "ex_*.txt",
                "mapping_sheet": "EXAMPLE_IN_Mapping",
                "target_staging_table": "STG_EXAMPLE",
                "thresholds_max_errors": None,  # blank -> default
            }
        ],
        output_files_rows=[
            {
                "file_type": "EXAMPLE_OUT",
                "glob": "ex_out_*.txt",
                "mapping_sheet": "EXAMPLE_OUT_Mapping",
                "rules_sheet": "EXAMPLE_OUT_Rules",
                "tolerance_max_errors": None,
                "tolerance_max_error_pct": None,
                "tolerance_ignore_fields": None,
            }
        ],
    )
    workbook = read_workbook(path)
    emitted_data = yaml.safe_load(emit_source_yaml(workbook))

    out_entry = emitted_data["output_files"][0]
    assert "strict_fixed_width" not in out_entry, (
        f"EA-S3 violation: default strict_fixed_width emitted in {out_entry!r}"
    )
    assert "strict_level" not in out_entry, (
        f"EA-S3 violation: default strict_level emitted in {out_entry!r}"
    )
    assert "tolerance" not in out_entry, (
        f"EA-S3 violation: default tolerance block emitted in {out_entry!r}"
    )

    in_entry = emitted_data["input_files"][0]
    assert "thresholds" not in in_entry, (
        f"EA-S3 violation: default thresholds block emitted in {in_entry!r}"
    )


# ---------------------------------------------------------------------------
# 3. Surgical strip preserves non-default tolerance.ignore_fields.
# ---------------------------------------------------------------------------


def test_emitted_yaml_preserves_non_default_tolerance_ignore_fields(tmp_path):
    """Surgical strip: ``ignore_fields`` override -> ``tolerance:`` block
    carries ONLY that sub-field; default ``max_errors`` / ``max_error_pct``
    siblings are dropped.

    Mirrors the SRC_A.yml pattern preserved verbatim by EA-S3.
    """
    path = _build_workbook(
        tmp_path,
        output_files_rows=[
            {
                "file_type": "EXAMPLE_OUT",
                "glob": "ex_*.txt",
                "mapping_sheet": "EXAMPLE_OUT_Mapping",
                "rules_sheet": "EXAMPLE_OUT_Rules",
                "tolerance_max_errors": None,
                "tolerance_max_error_pct": None,
                "tolerance_ignore_fields": "FILE_CREATE_TS",
            }
        ],
    )
    workbook = read_workbook(path)
    emitted_data = yaml.safe_load(emit_source_yaml(workbook))

    out_entry = emitted_data["output_files"][0]
    assert out_entry["tolerance"] == {"ignore_fields": ["FILE_CREATE_TS"]}, (
        f"Surgical strip lost or augmented: {out_entry['tolerance']!r}"
    )


# ---------------------------------------------------------------------------
# 4. Multi-record legacy keys never written (EB-S1 / ADR 0005).
# ---------------------------------------------------------------------------


def test_emitted_yaml_never_writes_multi_record_keys(tmp_path):
    """Mixed flat + umbrella output entries -- no ``multi_record`` or
    ``discriminator_field`` keys appear anywhere in the emitted YAML."""
    path = _build_workbook(
        tmp_path,
        output_files_rows=[
            {
                "file_type": "FLAT_OUT",
                "glob": "flat_*.txt",
                "mapping_sheet": "FLAT_OUT_Mapping",
                "rules_sheet": "FLAT_OUT_Rules",
                "tolerance_max_errors": None,
                "tolerance_max_error_pct": None,
                "tolerance_ignore_fields": None,
            },
            {
                "file_type": "UMBRELLA_OUT",
                "glob": "umbrella_*.txt",
                "mapping_sheet": "(umbrella)",
                "rules_sheet": "(umbrella)",
                "tolerance_max_errors": None,
                "tolerance_max_error_pct": None,
                "tolerance_ignore_fields": None,
            },
        ],
    )
    workbook = read_workbook(path)
    emitted = emit_source_yaml(workbook)

    # YAML-level key-form string check: the legacy keys would appear as
    # `multi_record:` / `discriminator_field:` lines. Plain substring
    # checks would false-positive on the legal `multi_record_report:`
    # gate name, so we check the key-form (with trailing colon) instead.
    # Restrict the search to lines inside output_files entries (i.e.
    # indented `- ` blocks before the gates: top-level key) to avoid
    # snagging any future legitimate use.
    output_section, _, _ = emitted.partition("\ngates:")
    assert "multi_record:" not in output_section, (
        f"EB-S1 violation: legacy 'multi_record:' key emitted.\n{output_section}"
    )
    assert "discriminator_field:" not in output_section, (
        f"EB-S1 violation: legacy 'discriminator_field:' key emitted.\n"
        f"{output_section}"
    )

    # Plus a structural check on every output entry (cannot false-
    # positive: dict keys are exact-match).
    emitted_data = yaml.safe_load(emitted)
    for entry in emitted_data["output_files"]:
        assert "multi_record" not in entry, (
            f"EB-S1 violation: 'multi_record' key in output entry {entry!r}"
        )
        assert "discriminator_field" not in entry, (
            f"EB-S1 violation: 'discriminator_field' key in output entry {entry!r}"
        )


# ---------------------------------------------------------------------------
# 5. Mapping path extension drives multi-record dispatch (EB-S1).
# ---------------------------------------------------------------------------


def test_emitted_yaml_mapping_path_extension_drives_multi_record(tmp_path):
    """``mapping_sheet`` = ``(umbrella)`` -> ``.yaml``; flat sheet -> ``.json``.

    Confirms the single-source-of-truth rule from ADR 0005: the file
    extension on ``mapping:`` is the only mechanism the engine uses to
    decide multi-record dispatch.
    """
    path = _build_workbook(
        tmp_path,
        output_files_rows=[
            {
                "file_type": "UMBRELLA",
                "glob": "umbrella_*.txt",
                "mapping_sheet": "(umbrella)",
                "rules_sheet": "(umbrella)",
                "tolerance_max_errors": None,
                "tolerance_max_error_pct": None,
                "tolerance_ignore_fields": None,
            },
            {
                "file_type": "FLAT",
                "glob": "flat_*.txt",
                "mapping_sheet": "FLAT_Mapping",
                "rules_sheet": "FLAT_Rules",
                "tolerance_max_errors": None,
                "tolerance_max_error_pct": None,
                "tolerance_ignore_fields": None,
            },
        ],
    )
    workbook = read_workbook(path)
    emitted_data = yaml.safe_load(emit_source_yaml(workbook))

    umbrella_entry = emitted_data["output_files"][0]
    flat_entry = emitted_data["output_files"][1]

    assert umbrella_entry["mapping"].endswith(".yaml"), (
        f"Umbrella mapping path should end with .yaml; got "
        f"{umbrella_entry['mapping']!r}"
    )
    assert flat_entry["mapping"].endswith(".json"), (
        f"Flat mapping path should end with .json; got "
        f"{flat_entry['mapping']!r}"
    )


# ---------------------------------------------------------------------------
# 6. Rules path is empty for umbrella outputs (per SHAW convention).
# ---------------------------------------------------------------------------


def test_emitted_yaml_rules_path_empty_for_umbrella(tmp_path):
    """``rules_sheet`` = ``(umbrella)`` -> emitted ``rules: ""``.

    Per-record rules live inside the umbrella mapping YAML; the
    output-file's ``rules:`` slot is intentionally empty to avoid
    double-declaration.
    """
    path = _build_workbook(
        tmp_path,
        output_files_rows=[
            {
                "file_type": "UMBRELLA",
                "glob": "umbrella_*.txt",
                "mapping_sheet": "(umbrella)",
                "rules_sheet": "(umbrella)",
                "tolerance_max_errors": None,
                "tolerance_max_error_pct": None,
                "tolerance_ignore_fields": None,
            },
        ],
    )
    workbook = read_workbook(path)
    emitted_data = yaml.safe_load(emit_source_yaml(workbook))

    assert emitted_data["output_files"][0]["rules"] == ""


# ---------------------------------------------------------------------------
# 7. Gates block emits all 7 gates with correct blocking + invoke_java.
# ---------------------------------------------------------------------------


def test_gates_block_emitted_from_source_sheet(tmp_path):
    """All 7 canonical gates appear in the emitted ``gates:`` block with
    the right ``blocking`` / ``invoke_java`` values.

    Only ``load_step`` and ``generate_step`` derive ``invoke_java`` from
    the workbook; the other 5 hardcode ``invoke_java: false``.
    """
    path = _build_workbook(
        tmp_path,
        source_overrides={
            "gate_load_blocking": "true",
            "gate_load_invoke_java": "true",
            "gate_f2s_blocking": "true",
            "gate_generate_blocking": "true",
            "gate_generate_invoke_java": "true",
            "gate_l1_blocking": "false",
            "gate_l2b_blocking": "true",
            "gate_l3_blocking": "false",
            "gate_mr_report_blocking": "false",
        },
    )
    workbook = read_workbook(path)
    emitted_data = yaml.safe_load(emit_source_yaml(workbook))

    gates = emitted_data["gates"]
    expected_names = {
        "load_step",
        "file_to_staging",
        "generate_step",
        "L1_structural",
        "L3_baseline_diff",
        "L2b_sql_truth",
        "multi_record_report",
    }
    assert set(gates) == expected_names, (
        f"Expected gates {expected_names}; got {set(gates)}"
    )

    # Per-gate value spot checks.
    assert gates["load_step"] == {"blocking": True, "invoke_java": True}
    assert gates["file_to_staging"] == {"blocking": True, "invoke_java": False}
    assert gates["generate_step"] == {"blocking": True, "invoke_java": True}
    assert gates["L1_structural"] == {"blocking": False, "invoke_java": False}
    assert gates["L2b_sql_truth"] == {"blocking": True, "invoke_java": False}
    assert gates["L3_baseline_diff"] == {"blocking": False, "invoke_java": False}
    assert gates["multi_record_report"] == {"blocking": False, "invoke_java": False}


# ---------------------------------------------------------------------------
# 8. Staging tables derived from input files, deduplicated, in input order.
# ---------------------------------------------------------------------------


def test_staging_tables_derived_from_input_files(tmp_path):
    """3 input files writing into 2 distinct staging tables -> the
    ``staging_tables:`` block carries both, deduplicated, in input-file
    order. Confirms the dedup keeps first-seen order."""
    path = _build_workbook(
        tmp_path,
        input_files_rows=[
            {
                "file_type": "FILE_A",
                "glob": "a_*.txt",
                "mapping_sheet": "FILE_A_Mapping",
                "target_staging_table": "STG_ALPHA",
                "thresholds_max_errors": None,
            },
            {
                "file_type": "FILE_B",
                "glob": "b_*.txt",
                "mapping_sheet": "FILE_B_Mapping",
                "target_staging_table": "STG_BETA",
                "thresholds_max_errors": None,
            },
            {
                "file_type": "FILE_C",
                "glob": "c_*.txt",
                "mapping_sheet": "FILE_C_Mapping",
                "target_staging_table": "STG_ALPHA",  # duplicate -> dropped
                "thresholds_max_errors": None,
            },
        ],
    )
    workbook = read_workbook(path)
    emitted_data = yaml.safe_load(emit_source_yaml(workbook))

    assert emitted_data["staging_tables"] == ["STG_ALPHA", "STG_BETA"], (
        f"Staging table dedup/order broken: {emitted_data['staging_tables']!r}"
    )


# ---------------------------------------------------------------------------
# 9. EmitterError on SourceConfig validation failure.
# ---------------------------------------------------------------------------


def test_emitter_validates_through_SourceConfig(tmp_path, monkeypatch):
    """A workbook that produces an invalid YAML raises :class:`EmitterError`,
    not a raw Pydantic ``ValidationError``.

    ``SourceConfig`` currently accepts blank ``source`` (``extra='allow'`` +
    no ``min_length`` constraint), so a real workbook cell can't trip the
    validator without extending the model -- which EC-S3 must not do
    (Sprint 1 schema is frozen). Instead we monkeypatch
    ``SourceConfig.model_validate`` to raise a synthetic
    ``ValidationError`` and confirm the emitter wraps it as
    ``EmitterError`` with the original exception chained as ``__cause__``.
    """
    from pydantic import BaseModel, Field

    # Construct a real synthetic ValidationError by validating a model
    # against bad input -- ValidationError has no public constructor,
    # so we bounce one through a Pydantic call.
    class _Probe(BaseModel):
        n: int = Field(gt=0)

    try:
        _Probe.model_validate({"n": -1})
    except Exception as raised:  # noqa: BLE001 -- intentional broad catch
        synthetic_error = raised

    # Patch SourceConfig.model_validate on the emitter's import site so
    # the emitter's own ``SourceConfig`` reference picks up the fake.
    from src.onboarding.emitters import source_yaml_emitter as emitter_mod

    def _always_raise(*_args, **_kwargs):
        raise synthetic_error

    monkeypatch.setattr(
        emitter_mod.SourceConfig, "model_validate", staticmethod(_always_raise)
    )

    # A perfectly-valid workbook -- the failure comes from the patched
    # validator, proving the emitter routes ValidationError -> EmitterError.
    path = _build_workbook(tmp_path)
    workbook = read_workbook(path)

    with pytest.raises(EmitterError) as exc_info:
        SourceYamlEmitter().emit(workbook)

    assert exc_info.value.__cause__ is synthetic_error
    assert "SourceConfig" in str(exc_info.value)


# ---------------------------------------------------------------------------
# 10. Emitted YAML passes the EB-S2 banned-keys guardrail.
# ---------------------------------------------------------------------------


def test_emitted_yaml_passes_eb_s2_guardrail(tmp_path):
    """Write the emitted SHAW YAML under a tmp ``sources/`` dir and run
    the EB-S2 per-file scanner against it; zero violations expected."""
    workbook = read_workbook(SHAW_WORKBOOK)
    emitted = emit_source_yaml(workbook)

    sources_dir = tmp_path / "sources"
    sources_dir.mkdir()
    target = sources_dir / "SHAW.yml"
    target.write_text(emitted, encoding="utf-8")

    # Reuse the EB-S2 scanner directly (no engine instantiation needed --
    # it is a pure-function helper in the guardrail test module).
    from tests.unit.test_source_yaml_guardrails import _scan_source_yaml

    violations = _scan_source_yaml(target)
    assert violations == [], (
        "EB-S2 guardrail flagged the emitted SHAW.yml:\n"
        + "\n".join(f"  - {v.location}: {v.key} = {v.value!r}" for v in violations)
    )
