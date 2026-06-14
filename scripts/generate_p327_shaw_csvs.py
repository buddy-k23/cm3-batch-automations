"""Generate mapping and rules CSVs from P327_SHAW_M06.xlsx.

Reads the TARGET section (columns N-Z) and Transformation Logic (column M)
from the p327 sheet and produces two Valdo-compatible CSV files:
  - mappings/csv/p327_shaw/P327_SHAW_M06_mapping.csv
  - mappings/csv/p327_shaw/P327_SHAW_M06_rules.csv
"""

import csv
import re
from pathlib import Path

import openpyxl

EXCEL_PATH = Path("mappings/excel/P327_SHAW_M06.xlsx")
OUT_DIR = Path("mappings/csv/p327_shaw")
OUT_DIR.mkdir(parents=True, exist_ok=True)

MAPPING_OUT = OUT_DIR / "P327_SHAW_M06_mapping.csv"
RULES_OUT = OUT_DIR / "P327_SHAW_M06_rules.csv"


def _clean(val, max_len=200):
    """Normalise a cell value to a single-line trimmed string."""
    if val is None:
        return ""
    s = str(val).strip().replace("\r", "").replace("\n", " ")
    # Collapse multiple spaces
    s = re.sub(r" {2,}", " ", s)
    return s[:max_len]


def _snake(name):
    """Convert FIELD-NAME to field_name."""
    return name.lower().replace("-", "_").replace(" ", "_")


# Field-specific valid value overrides — based on what SHAW actually sends
# (derived from Transformation Logic column, not generic APPS system values).
_VALID_VALUES_MAP = {
    # From Transformation Logic — SHAW-specific outputs
    "LOCATION-CODE": "100010|100020|100030|100040|100050|200000",
    "COLL-ASGN-UD-TYPE1": "0|1",
    "UD-7-TYPE": "I",
    "UD-3-TYPE": "1|2",
    "UD-5-FMT": "0|1|2",
    "USER-DEFINED-6": "DPD",
    "UD-6-FMT": "01|02|03|04|05|06|07",
    "PAYMENT-FREQUENCY": "MO|QT|SA|AN|BW|WK| ",
    "CUT-OFF-CODE": "CL|AC",
    "APP-CODE": "000",
    "NATL-CURRENCY": "USD",
    "BASE-CURRENCY": "USD",
    "RECOMMENDED-ACT-CD": "LS",
    "ACTG-SYS-ID": "LS|CAS|CASW",
    "ORIG-PORTF-TYPE": "Y|N",
    "CHRG-OFF-DOWN-IND": "0|1",
    "PRODUCT-TYPE": "R28|R31|R32|R33|R34|R36|C30|C31|R07|R10|R20|R25|R26|R27|R03|R02|R04|R16|R35|R17|R08|R13|R15|R19|R98|R99|NA",
}


def _extract_valid_values(field_name, raw_vv):
    """Extract pipe-separated valid value codes from raw spec text.

    Returns a pipe-separated string of codes, or empty string if none found.
    """
    # Use override map if available
    if field_name in _VALID_VALUES_MAP:
        return _VALID_VALUES_MAP[field_name]

    if not raw_vv:
        return ""

    # Skip purely descriptive text
    lower = raw_vv.lower()
    skip_starts = [
        "valid ", "must ", "see ", "refer ", "defined ", "the ", "this ",
        "used ", "for ", "if ", "when ", "each ", "cycle ",
        "an asterisk",
    ]
    if any(lower.startswith(p) for p in skip_starts):
        return ""

    # If already pipe-separated short codes, return as-is
    if "|" in raw_vv and len(raw_vv) < 80:
        return raw_vv

    return ""


def main():
    wb = openpyxl.load_workbook(str(EXCEL_PATH), data_only=True)
    ws = wb["p327"]

    # ── Extract fields ──────────────────────────────────────────────
    fields = []
    for row in ws.iter_rows(min_row=3, max_row=ws.max_row, values_only=True):
        fname = _clean(row[16], 60)  # col Q: Field Name
        if not fname:
            continue

        desc = _clean(row[17], 80)       # col R: Description
        pos_raw = _clean(row[19], 10)    # col T: Position
        len_raw = _clean(row[20], 10)    # col U: Length
        dtype = _clean(row[21], 20) or "String"  # col V: Data Type
        fmt = _clean(row[22], 30)        # col W: Format
        req = _clean(row[23], 5) or "N"  # col X: Required?
        vv = _clean(row[24], 120)        # col Y: Valid Values
        trans = _clean(row[12], 120)     # col M: Transformation Logic
        notes = _clean(row[25], 120)     # col Z: Notes

        # Fix position: take first integer if compound like "2-18"
        pos = ""
        if pos_raw:
            m = re.match(r"(\d+)", pos_raw)
            if m:
                pos = m.group(1)

        # Fix length: must be a plain integer
        length = ""
        if len_raw and len_raw.isdigit():
            length = len_raw
        else:
            # Derive from format if possible: 9(5)->5, +9(12)V9(6)->19, X(18)->18
            if fmt:
                nums = re.findall(r"\d+", fmt.replace("CCYYMMDD", ""))
                if "V" in fmt and len(nums) >= 2:
                    # e.g. +9(12)V9(6) -> 12+6+1(sign) = 19
                    length = str(int(nums[0]) + int(nums[1]) + (1 if fmt.startswith("+") or fmt.startswith("-") else 0))
                elif nums:
                    length = nums[0]
            if not length:
                # Fallback: estimate from dtype
                length = "19" if dtype == "Numeric" else ("8" if dtype == "Date" else "10")

        # Extract valid values: parse codes from descriptive text
        vv = _extract_valid_values(fname, vv)

        # Normalise required
        required = "Yes" if req.upper() in ("Y", "YES") else "No"

        fields.append({
            "field_name": fname,
            "data_type": dtype,
            "position": pos,
            "length": length,
            "target_name": _snake(fname),
            "required": required,
            "format": fmt,
            "transformation": trans,
            "valid_values": vv,
            "description": desc,
        })

    # ── Fix ESTIMATED-MNTH-PAY (bad length cell in Excel: '\xa0M-PO-AMOUNT') ──
    for f in fields:
        if f["field_name"] == "ESTIMATED-MNTH-PAY":
            f["length"] = "19"  # position 1954, next field ACI-INDICATOR at 1973

    # ── Write Mapping CSV ───────────────────────────────────────────
    mapping_cols = [
        "Field Name", "Data Type", "Position", "Length", "Target Name",
        "Required", "Format", "Transformation", "Valid Values", "Description",
    ]
    # Remap field dict keys to Title Case for the CSV header
    for f in fields:
        f["Field Name"] = f.pop("field_name")
        f["Data Type"] = f.pop("data_type")
        f["Position"] = f.pop("position")
        f["Length"] = f.pop("length")
        f["Target Name"] = f.pop("target_name")
        f["Required"] = f.pop("required")
        f["Format"] = f.pop("format")
        f["Transformation"] = f.pop("transformation")
        f["Valid Values"] = f.pop("valid_values")
        f["Description"] = f.pop("description")
    with open(MAPPING_OUT, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=mapping_cols, extrasaction="ignore")
        w.writeheader()
        for f in fields:
            w.writerow(f)

    print(f"Mapping CSV: {MAPPING_OUT} ({len(fields)} fields)")

    # ── Generate Rules CSV ──────────────────────────────────────────
    rules = []
    rule_id = 0
    cr_id = 0

    for f in fields:
        fname = f["Field Name"]
        dtype = f["Data Type"]
        fmt = f["Format"]
        req = f["Required"]
        length = f["Length"]
        vv = f["Valid Values"]

        # R: not_empty for required fields
        if req == "Yes":
            rule_id += 1
            rules.append({
                "Rule ID": f"R{rule_id:03d}",
                "Rule Name": f"{fname} required",
                "Field": fname,
                "Type": "not_empty",
                "Severity": "error",
                "Enabled": "Yes",
                "Message": f"{fname} must not be empty",
                "Value": "",
            })

        # R: exact_length for required fixed-width fields
        if req == "Yes" and length.isdigit():
            rule_id += 1
            rules.append({
                "Rule ID": f"R{rule_id:03d}",
                "Rule Name": f"{fname} length check",
                "Field": fname,
                "Type": "exact_length",
                "Severity": "error",
                "Enabled": "Yes",
                "Message": f"{fname} must be exactly {length} characters",
                "Value": length,
            })

        # R: numeric for Numeric fields
        if dtype == "Numeric":
            rule_id += 1
            rules.append({
                "Rule ID": f"R{rule_id:03d}",
                "Rule Name": f"{fname} numeric",
                "Field": fname,
                "Type": "numeric",
                "Severity": "error",
                "Enabled": "Yes",
                "Message": f"{fname} must be numeric",
                "Value": "",
            })

        # R: date_format for Date fields
        if dtype == "Date" and fmt:
            rule_id += 1
            rules.append({
                "Rule ID": f"R{rule_id:03d}",
                "Rule Name": f"{fname} date format",
                "Field": fname,
                "Type": "date_format",
                "Severity": "error",
                "Enabled": "Yes",
                "Message": f"{fname} must match {fmt}",
                "Value": fmt,
            })

        # R: valid_values if specific codes exist
        if vv and "|" in vv:
            rule_id += 1
            rules.append({
                "Rule ID": f"R{rule_id:03d}",
                "Rule Name": f"{fname} valid values",
                "Field": fname,
                "Type": "valid_values",
                "Severity": "error",
                "Enabled": "Yes",
                "Message": f"{fname} must be one of: {vv}",
                "Value": vv,
            })
        elif vv and len(vv) <= 20 and not " " in vv:
            # Single valid value like "USD"
            rule_id += 1
            rules.append({
                "Rule ID": f"R{rule_id:03d}",
                "Rule Name": f"{fname} valid values",
                "Field": fname,
                "Type": "valid_values",
                "Severity": "error",
                "Enabled": "Yes",
                "Message": f"{fname} must be: {vv}",
                "Value": vv,
            })

    # CR: LOCATION-CODE + ACCT-NUM unique composite
    cr_id += 1
    rules.append({
        "Rule ID": f"CR{cr_id:03d}",
        "Rule Name": "Account key unique",
        "Field": "LOCATION-CODE|ACCT-NUM",
        "Type": "cross_row:unique_composite",
        "Severity": "error",
        "Enabled": "Yes",
        "Message": "LOCATION-CODE + ACCT-NUM must be unique across all rows",
        "Value": "",
    })

    # ── Remap rule keys to BA-friendly column names ────────────────
    for r in rules:
        r["Rule Type"] = r.pop("Type")
        r["Expected / Values"] = r.pop("Value")

    rules_cols = [
        "Rule ID", "Rule Name", "Field", "Rule Type", "Severity",
        "Enabled", "Message", "Expected / Values",
    ]
    with open(RULES_OUT, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=rules_cols)
        w.writeheader()
        for r in rules:
            w.writerow(r)

    print(f"Rules CSV:   {RULES_OUT} ({len(rules)} rules)")


if __name__ == "__main__":
    main()
