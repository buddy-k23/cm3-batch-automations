"""Dump the Batch Header sheet raw for hand-authoring the CSV."""
import openpyxl
from pathlib import Path

wb = openpyxl.load_workbook(
    Path("mappings/excel/TRANERT_SHAW_Mappings.xlsx"), data_only=True
)
ws = wb["Batch Header"]
for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=9):
    vals = [str(c.value).strip() if c.value is not None else "" for c in row]
    print("|".join(vals))
