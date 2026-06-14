"""Show the full Valid Values list for NEW1 TRN-COD-ERT.

Diagnostic helper for TRANERT umbrella design. NEW1 covers multiple
transaction codes (32000, 32001, …) per the workbook spec; this
script extracts the complete list so the umbrella can wire one
record-type entry per code.
"""

from pathlib import Path
import openpyxl


def main() -> int:
    wb = openpyxl.load_workbook(
        Path("mappings/excel/TRANERT_SHAW_Mappings.xlsx"), data_only=True
    )
    ws = wb["NEW1"]
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=9):
        cell0 = row[0].value
        if cell0 and "TRN-COD" in str(cell0).upper():
            full_vv = row[7].value or ""
            print(f"Field: {cell0}")
            print(f"Position: {row[2].value}  Length: {row[3].value}")
            print(f"Valid Values cell (full):")
            for line in str(full_vv).splitlines():
                line = line.strip()
                if line:
                    print(f"  {line}")
            break
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
